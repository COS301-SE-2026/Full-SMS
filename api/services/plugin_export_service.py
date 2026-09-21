import io
import json
import csv
from typing import Any, Dict, Optional

try:
    import h5py

    HAS_H5PY = True
except ImportError:
    HAS_H5PY = False

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from api.services.plugin_service import get_execution_by_id

DATA_OUTPUT_TYPES = ["dataset", "array", "timeseries", "fitresult", "dataframe"]
PRESENTATION_OUTPUT_TYPES = ["plot", "histogram", "table", "value", "heatmap"]


def get_output_config(execution: dict, output_id: str) -> Optional[dict]:
    plugin_info = execution.get("user_plugins", {})
    config = plugin_info.get("config", {})
    outputs = config.get("outputs", [])

    for output in outputs:
        if output.get("id") == output_id:
            return output
    return None


def _csv_write_array(writer, data: Any) -> None:
    values = data.get("values", data) if isinstance(data, dict) else data
    writer.writerow(["index", "value"])
    for i, val in enumerate(values):
        writer.writerow([i, val])


def _csv_write_timeseries(writer, data: dict) -> None:
    time = data.get("time", [])
    values = data.get("values", [])
    time_unit = data.get("timeUnit", "s")
    writer.writerow([f"time ({time_unit})", "value"])
    for t, v in zip(time, values):
        writer.writerow([t, v])


def _csv_write_dataframe(writer, data: dict) -> None:
    columns = data.get("columns", [])
    if not columns:
        return
    headers = [col.get("name", f"col_{i}") for i, col in enumerate(columns)]
    writer.writerow(headers)
    num_rows = len(columns[0].get("values", [])) if columns else 0
    for row_idx in range(num_rows):
        row = [col.get("values", [])[row_idx] for col in columns]
        writer.writerow(row)


def _csv_write_table(writer, data: dict) -> None:
    col_headers = data.get("columns", [])
    rows = data.get("rows", [])
    if col_headers:
        writer.writerow(col_headers)
    for row in rows:
        writer.writerow(row)


def _csv_write_fallback(writer, data: Any) -> None:
    if isinstance(data, list):
        writer.writerow(["index", "value"])
        for i, val in enumerate(data):
            writer.writerow([i, val])
    elif isinstance(data, dict):
        writer.writerow(["key", "value"])
        for k, v in data.items():
            writer.writerow([k, v])


def export_to_csv(data: Any, output_type: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    if output_type == "array":
        _csv_write_array(writer, data)
    elif output_type == "timeseries":
        _csv_write_timeseries(writer, data)
    elif output_type == "dataframe":
        _csv_write_dataframe(writer, data)
    elif output_type == "table":
        _csv_write_table(writer, data)
    else:
        _csv_write_fallback(writer, data)

    return buffer.getvalue().encode("utf-8")


def export_to_json(data: Any, _output_type: str) -> bytes:
    return json.dumps(data, indent=2, default=str).encode("utf-8")


def _excel_write_array(ws, data: Any) -> None:
    values = data.get("values", data) if isinstance(data, dict) else data
    ws.append(["index", "value"])
    for i, val in enumerate(values):
        ws.append([i, val])


def _excel_write_timeseries(ws, data: dict) -> None:
    time = data.get("time", [])
    values = data.get("values", [])
    time_unit = data.get("timeUnit", "s")
    ws.append([f"time ({time_unit})", "value"])
    for t, v in zip(time, values):
        ws.append([t, v])


def _excel_write_dataframe(ws, data: dict) -> None:
    columns = data.get("columns", [])
    if not columns:
        return
    headers = [col.get("name", f"col_{i}") for i, col in enumerate(columns)]
    ws.append(headers)
    num_rows = len(columns[0].get("values", [])) if columns else 0
    for row_idx in range(num_rows):
        row = [col.get("values", [])[row_idx] for col in columns]
        ws.append(row)


def _excel_write_table(ws, data: dict) -> None:
    col_headers = data.get("columns", [])
    rows = data.get("rows", [])
    if col_headers:
        ws.append(col_headers)
    for row in rows:
        ws.append(list(row))


def _excel_write_fitresult(wb, ws, data: dict) -> None:
    ws.append(["Parameter", "Value", "Uncertainty"])
    params = data.get("parameters", {})
    uncertainties = data.get("uncertainties", {})
    for key, value in params.items():
        ws.append([key, value, uncertainties.get(key, "N/A")])

    ws_metrics = wb.create_sheet("Metrics")
    ws_metrics.append(["Metric", "Value"])
    for key, value in data.get("metrics", {}).items():
        ws_metrics.append([key, value])


def _excel_write_dataset(wb, ws, data: dict) -> None:
    arrays = data.get("arrays", {})
    first = True
    for name, arr_data in arrays.items():
        if first:
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        ws.append(["index", "value"])
        for i, val in enumerate(arr_data.get("values", [])):
            ws.append([i, val])


def _excel_write_fallback(ws, data: Any) -> None:
    if isinstance(data, dict):
        ws.append(["key", "value"])
        for k, v in data.items():
            ws.append([k, str(v)])
    elif isinstance(data, list):
        ws.append(["index", "value"])
        for i, val in enumerate(data):
            ws.append([i, val])


def export_to_excel(data: Any, output_type: str) -> bytes:
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl is required for Excel export")

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Plugin Output"

    if output_type == "array":
        _excel_write_array(ws, data)
    elif output_type == "timeseries":
        _excel_write_timeseries(ws, data)
    elif output_type == "dataframe":
        _excel_write_dataframe(ws, data)
    elif output_type == "table":
        _excel_write_table(ws, data)
    elif output_type == "fitresult":
        _excel_write_fitresult(wb, ws, data)
    elif output_type == "dataset":
        _excel_write_dataset(wb, ws, data)
    else:
        _excel_write_fallback(ws, data)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _hdf5_write_array(f, data: Any) -> None:
    import numpy as np

    values = data.get("values", data) if isinstance(data, dict) else data
    f.create_dataset("values", data=np.array(values))
    if not isinstance(data, dict):
        return
    for attr in ["unit", "label", "dtype"]:
        if data.get(attr):
            f["values"].attrs[attr] = data[attr]


def _hdf5_write_timeseries(f, data: dict) -> None:
    import numpy as np

    f.create_dataset("time", data=np.array(data.get("time", [])))
    f.create_dataset("values", data=np.array(data.get("values", [])))
    f["time"].attrs["unit"] = data.get("timeUnit", "s")
    if data.get("valueUnit"):
        f["values"].attrs["unit"] = data["valueUnit"]


def _hdf5_write_dataset_array(f, name: str, arr_data: dict) -> None:
    import numpy as np

    ds = f.create_dataset(name, data=np.array(arr_data.get("values", [])))
    for attr in ["unit", "dtype", "description"]:
        if arr_data.get(attr):
            ds.attrs[attr] = arr_data[attr]


def _hdf5_write_dataset(f, data: dict) -> None:
    for name, arr_data in data.get("arrays", {}).items():
        _hdf5_write_dataset_array(f, name, arr_data)
    if not data.get("metadata"):
        return
    meta_grp = f.create_group("metadata")
    for key, value in data["metadata"].items():
        if isinstance(value, (str, int, float, bool)):
            meta_grp.attrs[key] = value


def _hdf5_write_fitresult(f, data: dict) -> None:
    import numpy as np

    params_grp = f.create_group("parameters")
    for key, value in data.get("parameters", {}).items():
        params_grp.attrs[key] = value

    if data.get("uncertainties"):
        unc_grp = f.create_group("uncertainties")
        for key, value in data["uncertainties"].items():
            unc_grp.attrs[key] = value

    metrics_grp = f.create_group("metrics")
    for key, value in data.get("metrics", {}).items():
        if value is not None:
            metrics_grp.attrs[key] = value

    if data.get("fitted_curve"):
        curve = data["fitted_curve"]
        f.create_dataset("fitted_curve/x", data=np.array(curve.get("x", [])))
        f.create_dataset("fitted_curve/y", data=np.array(curve.get("y", [])))

    if data.get("residuals"):
        f.create_dataset("residuals", data=np.array(data["residuals"]))


def _hdf5_write_dataframe(f, data: dict) -> None:
    import numpy as np

    for col in data.get("columns", []):
        name = col.get("name", "unnamed")
        values = col.get("values", [])
        dtype = col.get("dtype", "float64")

        if dtype == "string":
            dt = h5py.special_dtype(vlen=str)
            ds = f.create_dataset(name, data=values, dtype=dt)
        else:
            ds = f.create_dataset(name, data=np.array(values))
        ds.attrs["dtype"] = dtype


HDF5_WRITERS = {
    "array": _hdf5_write_array,
    "timeseries": _hdf5_write_timeseries,
    "dataset": _hdf5_write_dataset,
    "fitresult": _hdf5_write_fitresult,
    "dataframe": _hdf5_write_dataframe,
}


def export_to_hdf5(data: Any, output_type: str) -> bytes:
    if not HAS_H5PY:
        raise ImportError("h5py is required for HDF5 export")

    buffer = io.BytesIO()

    with h5py.File(buffer, "w") as f:
        writer = HDF5_WRITERS.get(output_type)
        if writer:
            writer(f, data)
        else:
            f.attrs["data"] = json.dumps(data, default=str)

    buffer.seek(0)
    return buffer.getvalue()


def export_plugin_output(
    execution_id: str, output_id: str, format: str, user_id: str
) -> Dict[str, Any]:
    execution = get_execution_by_id(execution_id)

    if not execution:
        raise ValueError("Execution not found")

    plugin_info = execution.get("user_plugins", {})
    if plugin_info.get("user_id") != user_id:
        raise ValueError("Access denied")

    output_config = get_output_config(execution, output_id)
    if not output_config:
        raise ValueError(f"Output '{output_id}' not found in plugin config")

    if output_config.get("exportable") is False:
        raise ValueError("This output is not exportable")

    results = execution.get("results", {})
    if output_id not in results:
        raise ValueError(f"Output '{output_id}' not found in execution results")

    data = results[output_id]
    output_type = output_config.get("type", "value")

    valid_formats = get_valid_formats_for_type(output_type)
    if format not in valid_formats:
        raise ValueError(
            f"Format '{format}' is not supported for output type '{output_type}'"
        )

    if format == "csv":
        content = export_to_csv(data, output_type)
        content_type = "text/csv"
        extension = "csv"
    elif format == "json":
        content = export_to_json(data, output_type)
        content_type = "application/json"
        extension = "json"
    elif format == "excel":
        content = export_to_excel(data, output_type)
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        extension = "xlsx"
    elif format == "hdf5":
        content = export_to_hdf5(data, output_type)
        content_type = "application/x-hdf5"
        extension = "h5"
    else:
        raise ValueError(f"Unsupported export format: {format}")

    plugin_name = plugin_info.get("name", "plugin").replace(" ", "_")
    filename = f"{plugin_name}_{output_id}.{extension}"

    return {
        "content": content,
        "content_type": content_type,
        "filename": filename,
    }


def get_valid_formats_for_type(output_type: str) -> list:
    format_map = {
        "plot": ["png", "pdf", "svg", "json"],
        "histogram": ["png", "pdf", "svg", "json"],
        "heatmap": ["png", "pdf", "svg", "json"],
        "table": ["csv", "json", "excel"],
        "dataframe": ["csv", "json", "excel"],
        "value": ["json"],
        "array": ["csv", "json", "hdf5"],
        "timeseries": ["csv", "json", "hdf5"],
        "dataset": ["json", "hdf5"],
        "fitresult": ["json", "hdf5"],
    }
    return format_map.get(output_type, ["json"])


def export_all_outputs(execution_id: str, format: str, user_id: str) -> Dict[str, Any]:
    execution = get_execution_by_id(execution_id)

    if not execution:
        raise ValueError("Execution not found")

    plugin_info = execution.get("user_plugins", {})
    if plugin_info.get("user_id") != user_id:
        raise ValueError("Access denied")

    results = execution.get("results", {})
    config = plugin_info.get("config", {})
    outputs = config.get("outputs", [])

    export_data = {}
    for output in outputs:
        output_id = output.get("id")
        if output_id not in results:
            continue
        if output.get("exportable") is False:
            continue

        output_type = output.get("type", "value")
        valid_formats = get_valid_formats_for_type(output_type)

        if format in valid_formats:
            export_data[output_id] = {
                "label": output.get("label"),
                "type": output_type,
                "data": results[output_id],
            }

    if format == "json":
        content = json.dumps(export_data, indent=2, default=str).encode("utf-8")
        content_type = "application/json"
        extension = "json"
    else:
        raise ValueError("Bulk export only supports JSON format currently")

    plugin_name = plugin_info.get("name", "plugin").replace(" ", "_")
    filename = f"{plugin_name}_all_outputs.{extension}"

    return {
        "content": content,
        "content_type": content_type,
        "filename": filename,
    }
